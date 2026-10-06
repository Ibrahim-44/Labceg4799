from pathlib import Path
from datetime import date
import subprocess
from docx import Document
from docx.shared import Inches, Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'Rapport_Lab2_Amel_Ibrahim_Mat.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2)
sec.left_margin = sec.right_margin = Cm(2.2)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
for name in ('Heading 1', 'Heading 2', 'Heading 3'):
    doc.styles[name].font.color.rgb = RGBColor.from_string('16324F')
    doc.styles[name].paragraph_format.keep_with_next = True
code_style = doc.styles.add_style('Code laboratoire', 1)
code_style.font.name = 'Consolas'
code_style.font.size = Pt(8)
code_style.paragraph_format.space_after = Pt(0)
code_style.paragraph_format.line_spacing = 1
footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('Amel • Ibrahim Mat | Laboratoire 2 | ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def p(text):
    doc.add_paragraph(text)

def h(text, level=1):
    doc.add_heading(text, level)

def code(text):
    for line in text.strip('\n').splitlines():
        doc.add_paragraph(line, 'Code laboratoire')

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Shading Accent 1'
    for c, value in zip(t.rows[0].cells, headers):
        c.text = value
    for row in rows:
        cells = t.add_row().cells
        for c, value in zip(cells, row):
            c.text = str(value)
        trPr = t.rows[-1]._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
    doc.add_paragraph()

fig_no = 0
def figure(name, caption):
    global fig_no
    fig_no += 1
    w, ht = Image.open(ROOT / name).size
    width = min(6.45, 7.7 * w / ht)
    par = doc.add_paragraph()
    par.alignment = 1
    par.add_run().add_picture(str(ROOT / name), width=Inches(width))
    cap = doc.add_paragraph(f'Figure {fig_no} — {caption} ({name}).', 'Caption')
    cap.alignment = 1

doc.add_heading('Rapport du laboratoire 2', 0)
p('CEG 4799 A00 / CSI 4539 A00\nConception de systèmes informatiques sécuritaires\nAutomne 2026')
doc.add_heading('Attaques mémoire, outils de bas niveau et défense par le typage', 1)
p('Équipe : Amel et Ibrahim Mat')
p('Professeur : Mohamed Ali Ibrahim\nAssistant à l’enseignement : Mohamed Nefsi')
p('Rapport préparé le 5 octobre 2026\nÉnoncé de référence : version 1.1 du 3 octobre 2026')
p('Ce rapport suit les neuf rubriques de la section 7.1 de l’énoncé. Les résultats sont établis à partir des captures et fichiers fournis. Les essais manquants ou incomplets sont indiqués explicitement; les comportements attendus ne sont pas présentés comme des résultats obtenus.')
table(['Document', 'Format'], [('Version de travail demandée', 'Word (.docx), modifiable'), ('Format de remise prescrit par l’énoncé', 'PDF, à exporter depuis Word après finalisation')])
doc.add_page_break()
h('Plan du rapport')
for text in [
    '1. Introduction et modèle de menace',
    '2. Disposition de la pile et contre-mesures (E1, E4)',
    '3. Exploitations mémoire (E2, E3, E5)',
    '4. Outils de bas niveau — composante L2-B (E6)',
    '5. Défense par le typage et propriétés (E7, E8)',
    '6. Corpus légué et limites',
    '7. Extension à distance — version Serveur (E9)',
    '8. Travail d’équipe',
    '9. Sources et déclaration d’usage de l’IA',
    'Annexe A. Tableau des essais T1 à T8',
    'Annexe B. Code fourni et relevés du binaire',
    'Annexe C. Éléments à compléter pour la remise',
]:
    p(text)

h('1. Introduction et modèle de menace')
p('Le laboratoire étudie les erreurs de gestion de la mémoire en C, leur effet sur le flot de contrôle et les protections qui limitent leur exploitation. Il relie le code source, le désassemblage et les octets machine, puis examine une défense en Ada fondée sur des contraintes de longueur. Les essais documentés portent sur les programmes locaux du laboratoire; la partie Serveur est évaluée selon les preuves disponibles.')
h('1.1 Programme vulnérable et cause première', 2)
p('Dans stack.c, main lit jusqu’à 517 octets du fichier badfile dans str[517], puis appelle bof(str). La fonction bof réserve buffer[100] et appelle strcpy(buffer, str) sans vérifier la longueur. Une chaîne dépassant la capacité du tampon peut écraser les données voisines de la pile, dont l’adresse de retour.')
code('char buffer[100];\nstrcpy(buffer, str);')
p('La propriété minimale pour cette copie est strlen(str) + 1 <= sizeof(buffer), soit au plus 99 caractères utiles pour une chaîne C de capacité 100. Le caractère terminal nul compte dans l’espace nécessaire. De plus, fread ne garantit pas de terminaison nulle : si aucun nul n’est présent dans les octets lus, strcpy peut lire au-delà de str. La variable length est calculée mais n’est pas utilisée pour contrôler cette copie.')
h('1.2 Actifs, adversaire et surfaces d’entrée', 2)
table(['Élément', 'Description'], [
    ('Actifs protégés', 'Intégrité de la pile, adresse de retour, flot de contrôle, identité effective du processus et disponibilité du programme.'),
    ('Capacité locale de l’adversaire', 'Contrôler le contenu de badfile lu par le programme vulnérable.'),
    ('Surface locale', 'Lecture fread puis copie strcpy dans le tampon de taille fixe.'),
    ('Surface distante attendue', 'Données reçues par TCP sur le port 9090 du serveur SEED; aucune exploitation distante n’est visible dans les pièces fournies.'),
    ('Frontière de confiance', 'Octets externes vers mémoire du processus; dans le scénario Set-UID, code contrôlé par l’utilisateur vers privilèges du programme.'),
])
p('Une segmentation fault prouve un échec d’exécution, sans suffire à prouver un shell ni une élévation de privilèges. Pour établir le succès Set-UID, une trace de shell accompagnée de id ou getresuid est nécessaire. Aucune capture fournie ne présente cette preuve.')

h('2. Disposition de la pile et contre-mesures (E1, E4)')
h('2.1 Observation GDB et calcul du décalage', 2)
p('La capture E1 montre un SIGSEGV avec le compteur d’instruction à 0x90909090. Cette valeur correspond à quatre octets 0x90, utilisés comme NOP dans la charge. Elle est compatible avec un écrasement de la cible de retour par le remplissage. La commande p &buffer affiche 0xf7fb8ebc. La capture ne montre pas $ebp ni les réglages de compilation; elle ne permet donc pas, à elle seule, de calculer la distance exacte.')
figure('lab2e1.png', 'Arrêt sous GDB et adresse du tampon affichée')
p('Le script visible dans lab2e2.png suppose un décalage de 104 octets, calculé comme 100 + 4. Cette formule néglige les registres sauvegardés et l’alignement que le compilateur peut ajouter. Le binaire stack fourni a été désassemblé pendant la préparation du rapport : son tampon est adressé par [ebp-0x6c], soit EBP − 108, et l’adresse de retour est à EBP + 4. Le décalage de ce binaire est donc 108 + 4 = 112 octets. Ce résultat concerne stack fourni, et ne certifie pas la disposition d’un autre binaire, notamment stack-L1 ou retlib.')
code('Binaire stack fourni (32 bits), adresses croissantes :\nEBP - 108   début de buffer[100]\nEBP -   9   dernier octet de buffer\nEBP -   8   espace intercalaire (4 octets)\nEBP -   4   EBX sauvegardé (4 octets)\nEBP         EBP sauvegardé (4 octets)\nEBP +   4   adresse de retour (4 octets)\nEBP +   8   argument str\n\nDistance(buffer, retour) = (EBP + 4) - (EBP - 108) = 112')
p('Dans ce binaire, l’offset 104 correspond au début d’EBX sauvegardé. La disposition explique pourquoi une taille de tampon ne suffit pas à déduire l’offset de retour. Les adresses absolues varient selon le processus, l’ASLR et l’environnement; elles doivent être relevées pour le binaire effectivement testé.')
h('2.2 Résultats des contre-mesures', 2)
figure('lab2e4.png', 'ASLR actif, script de force brute absent et détection du canari')
p('La lecture de /proc/sys/kernel/randomize_va_space renvoie 2 : l’ASLR est actif au moment de cette capture. L’appel à ./brute-force.sh échoue avec « No such file or directory »; aucune force brute réussie n’est démontrée. La compilation gcc -m32 -z noexecstack -o stack-protected stack.c est suivie d’une exécution arrêtée par « stack smashing detected » puis « Aborted (core dumped) ». Cette trace établit la détection d’un canari endommagé dans cet essai. Comme NX et le protecteur de pile sont présents ensemble, elle n’isole pas l’effet de NX.')
table(['Protection', 'Mécanisme et limite', 'État des preuves'], [
    ('dash', 'Peut abandonner des privilèges lorsque les identités réelle et effective diffèrent. Le scénario demandé ajoute setuid(0) avant le shell; cela ne corrige pas la copie vulnérable.', 'Aucun lien /bin/sh ni essai setuid(0) montré. Le shellcode visible n’inclut pas cet appel.'),
    ('ASLR', 'Randomise les adresses et rend les valeurs absolues moins fiables. La force brute 32 bits reste une limite à étudier expérimentalement.', 'Valeur 2 visible. Script de force brute absent; contournement non démontré.'),
    ('StackGuard', 'Place un canari et contrôle son intégrité avant le retour. Ne vérifie pas toutes les écritures et ne supprime pas la cause de l’erreur.', 'Message de détection visible; lecture et contrôle du canari visibles dans la capture E6.'),
    ('NX', 'Interdit l’exécution du code placé sur la pile. Ne bloque pas, à lui seul, la réutilisation de code existant par return-to-libc.', '-z noexecstack visible dans un essai combiné; absence d’essai NX seul. Le binaire stack fourni a une pile rwx.'),
])
p('Pour une comparaison avant/après reproductible, il manque des traces où une seule protection change à la fois. Le réglage ASLR=2 observé en E4 ne doit pas être attribué aux autres essais sans preuve. De même, la présence du canari dépend du binaire et des options du compilateur; elle doit être vérifiée, plutôt que supposée à partir d’une option omise.')

h('3. Exploitations mémoire (E2, E3, E5)')
h('3.1 Shellcode 32 bits (E2)', 2)
p('La capture lab2e2.png contient une version de exploit.py qui place un shellcode de 23 octets à la fin d’un tableau de 517 NOP. Ce shellcode construit la chaîne /bin//sh sur la pile et prépare l’appel système execve avec int 0x80. Les octets ci-dessous sont transcrits de cette capture; ils ne sont pas ceux du script Python actuellement fourni, qui prépare une tentative de return-to-libc.')
code('31 c0 50 68 2f 2f 73 68 68 2f 62 69 6e 89 e3 50\n53 89 e1 b0 0b cd 80')
table(['Octets', 'Instruction', 'Rôle'], [
    ('31 c0', 'xor eax,eax', 'Met EAX à zéro sans immédiat nul.'),
    ('50', 'push eax', 'Ajoute le terminal nul au moment de l’exécution.'),
    ('68 2f 2f 73 68', 'push 0x68732f2f', 'Empile //sh en ordre little-endian.'),
    ('68 2f 62 69 6e', 'push 0x6e69622f', 'Empile /bin.'),
    ('89 e3', 'mov ebx,esp', 'Fait pointer EBX sur la chaîne.'),
    ('50 53 89 e1', 'push eax; push ebx; mov ecx,esp', 'Construit argv = {pointeur de chaîne, NULL}.'),
    ('b0 0b cd 80', 'mov al,0x0b; int 0x80', 'Appel execve selon l’ABI Linux x86 32 bits.'),
])
p('La séquence transcrite ne contient aucun octet 00. Elle initialise EAX, EBX et ECX, mais ne met pas explicitement EDX à zéro pour envp : la validité de ce registre au moment de l’appel n’est pas établie par la capture, ce qui limite la fiabilité du shellcode. Cela ne remplace pas les tests autonomes demandés avec a32.out et a64.out, ni l’assemblage avec nasm et ld. Les fichiers hello.s, mysh64.s, another_sh64.s, convert.py et call_shellcode.c ne sont pas présents dans le dossier fourni.')
p('Deux approches sont demandées par E2. Dans l’approche call, une instruction call empile l’adresse des données situées après elle; le code récupère cette adresse pour localiser la chaîne intégrée au segment de code. Dans l’approche sur la pile, les caractères sont empilés et ESP fournit l’adresse de la chaîne. Le shellcode visible utilise la seconde approche. Aucun développement ni test de shellcode 64 bits n’est documenté ici.')
figure('lab2e2.png', 'Version du script contenant le shellcode 32 bits et les paramètres de la charge')
h('3.2 Débordement de tampon (E3)', 2)
table(['Paramètre de la capture', 'Valeur', 'Analyse'], [
    ('Taille de content', '517 octets', 'Même taille que la lecture maximale de main dans le source fourni.'),
    ('Remplissage', '0x90', 'Zone de NOP destinée à augmenter la plage de cibles possibles.'),
    ('start', '517 − 23 = 494', 'Place le shellcode transcrit à partir de l’octet 494.'),
    ('ret', '0xffffc4bc + 100 = 0xffffc520', 'Valeur écrite dans la capture; elle n’est pas confirmée par le relevé GDB E1.'),
    ('offset', '104', 'Hypothèse de la capture; le binaire stack fourni nécessite 112 pour la cible de retour.'),
    ('Encodage', '4 octets little-endian', 'Cohérent avec un processus x86 32 bits.'),
])
figure('lab2e3.png', 'Génération de badfile puis échec de stack-L1')
p('Le fichier badfile est généré et ls -l affiche 517 octets. L’exécution ./stack-L1 se termine par « Segmentation fault ». Le verdict de T3 est donc un essai de débordement en échec pour l’objectif de shell root. Les causes possibles comprennent un offset incorrect, une adresse non valable ou des protections actives; cette capture ne permet pas de choisir une cause unique.')
p('Cause première : une copie sans contrôle de longueur dépasse le tampon. Propriétés violées : les écritures doivent rester dans buffer[0..99], et l’adresse de retour doit rester identique à sa valeur légitime jusqu’au retour de bof. Le crash ne démontre pas l’exécution de execve.')
p('Le Niveau 2 exige une charge unique pour une taille de tampon comprise entre 100 et 200 octets, sans force brute. Le Niveau 3 exige une exploitation 64 bits tenant compte des octets nuls des adresses. Aucun script ni résultat fourni ne démontre ces deux niveaux. Le Niveau 4 est un bonus, non documenté.')
h('3.3 Tentative de return-to-libc (E5)', 2)
figure('lab2e5.png', 'Adresses de system et exit relevées sous GDB dans retlib')
p('Sous GDB, un point d’arrêt est placé dans bof de retlib. Les commandes print system et print exit affichent respectivement 0xf7d768e0 et 0xf7d655b0. Ces adresses correspondent au processus observé; elles ne sont pas des constantes universelles. La capture ne montre ni l’adresse validée de /bin/sh, ni l’exécution de la charge, ni l’état NX ou l’identité effective.')
table(['Élément du script actuel', 'Valeur', 'Interprétation'], [
    ('Taille', '300 octets', 'Diffère de la charge de 517 octets visible pour le débordement.'),
    ('X : cible de retour', '104', 'system_addr est écrit à cet offset supposé.'),
    ('Y : retour de system', '108', 'exit_addr est écrit à X + 4.'),
    ('Z : argument de system', '112', 'sh_addr est écrit à X + 8.'),
    ('Chaîne intégrée', 'offset 200 : /bin/sh\\x00', 'Chaîne ajoutée directement dans badfile.'),
    ('Adresse de chaîne supposée', '0xf7f54ebc + 200 = 0xf7f54f84', 'Base non validée par la capture E1, qui affiche 0xf7fb8ebc.'),
])
p('Selon la convention d’appel x86 32 bits, un retour détourné vers system doit lui laisser une adresse de retour puis un argument pointant vers une chaîne terminée par zéro. Si ces données sont correctes, system peut exécuter du code déjà présent dans libc sans exécuter de code injecté sur la pile. Cela explique le contournement conceptuel de NX, mais aucune réussite n’est établie par les fichiers présents.')
p('Le script fourni diffère aussi de la procédure demandée : E5 prévoit une chaîne d’environnement exportée et son adresse obtenue par prtenv, tandis que le script intègre une chaîne dans la charge. Dans le cas particulier de strcpy, une adresse contenant un octet nul interrompt la copie : ici les adresses écrites n’ont pas d’octet nul, mais le terminal de /bin/sh interrompt normalement la copie après cette chaîne. retlib.c manque; il n’est donc pas possible de vérifier si ce raisonnement de copie s’applique à ce programme.')
p('La présence de exit n’est pas nécessaire pour que system commence l’exécution de la commande. Elle donne une destination de retour cohérente après system; sans retour valide, un crash peut survenir à la fermeture du shell. Son statut de sortie dépend de la valeur alors lue sur la pile et n’est pas contrôlé explicitement par ce script. Un changement de nom du binaire peut déplacer les données d’arguments et d’environnement, donc l’adresse d’une chaîne qui y est stockée. Ces deux variations sont expliquées, mais n’ont pas été testées dans les traces fournies.')
p('Cause première : l’intégrité de la cible de retour n’est pas préservée. Propriété violée : la cible doit rester le site de retour légitime, sans substitution par system. Verdict T5 : adresses relevées et charge préparée, exploitation non démontrée.')

h('4. Outils de bas niveau — composante L2-B (E6)')
h('4.1 Chaîne C–assembleur–code machine du binaire fourni', 2)
p('Une analyse supplémentaire, en lecture seule, a été effectuée sur le fichier stack fourni avec objdump -Mintel -d --disassemble=bof stack. Il s’agit d’un binaire ELF 32 bits i386. Ce relevé est distinct de la capture lab2e6.png, qui concerne retlib-protected.')
table(['Adresse / octets', 'Instruction', 'Lien avec le source et effet'], [
    ('08049a65 / 55', 'push ebp', 'Début du prologue : sauvegarde du pointeur de cadre.'),
    ('08049a66 / 89 e5', 'mov ebp,esp', 'Installe le cadre courant.'),
    ('08049a68 / 53', 'push ebx', 'Sauvegarde EBX; occupe de l’espace entre tampon et EBP.'),
    ('08049a69 / 83 ec 74', 'sub esp,0x74', 'Réserve l’espace du cadre; la réservation ne vaut pas seulement 100 octets.'),
    ('08049a79 / ff 75 08', 'push DWORD PTR [ebp+0x8]', 'Place l’argument source str sur la pile.'),
    ('08049a7c / 8d 55 94', 'lea edx,[ebp-0x6c]', 'Calcule l’adresse de buffer; base utilisée pour le calcul de 112 octets.'),
    ('08049a7f / 52', 'push edx', 'Place l’adresse du tampon destination.'),
    ('08049a82 / e8 99 f5 ff ff', 'call 0x8049020', 'Appel à la routine de copie correspondant à strcpy du source; cible affichée par objdump comme _init+0x20.'),
    ('08049a8a / b8 01 00 00 00', 'mov eax,0x1', 'Valeur de retour 1 de bof.'),
    ('08049a8f / 8b 5d fc', 'mov ebx,[ebp-0x4]', 'Restaure EBX sauvegardé.'),
    ('08049a92 / c9', 'leave', 'Épilogue : restaure ESP et EBP.'),
    ('08049a93 / c3', 'ret', 'Charge l’adresse de retour depuis la pile.'),
])
p('Les appels __x86.get_pc_thunk.ax et l’ajustement d’EAX participent à l’adressage des données du binaire. La fonction examinée ne contient pas le chargement et le contrôle de canari visibles dans retlib-protected. Le lien entre assembleur et machine du shellcode est présenté en section 3.1; aucun relevé NASM autonome n’a été fourni.')
h('4.2 Canari dans la capture de retlib-protected', 2)
figure('lab2e6.png', 'Désassemblage de bof avec lecture et vérification du canari')
p('Le prologue commence par push ebp (55), mov ebp,esp (89 e5), push ebx (53) puis sub esp,0x84. L’instruction à 0x11f7 lit DWORD PTR gs:0x14; à 0x11fe, la valeur est stockée à [ebp-0xc]. À 0x1220, elle est relue, puis comparée à gs:0x14 par sub. Si elle a changé, le chemin mène à __stack_chk_fail_local. Cette séquence explique le message « stack smashing detected » de la capture E4.')
p('La capture de retlib-protected montre un appel à fread avec une taille 0x12c (300) et une destination [ebp-0x70]. Elle ne montre pas l’appel strcpy de stack.c. Sans retlib.c, il serait incorrect d’attribuer ce désassemblage au source stack.c fourni. La capture s’arrête avant leave et ret; ces instructions sont identifiées dans le relevé complet de stack, en section 4.1 et en annexe B.')
h('4.3 Effet des options de compilation et limites', 2)
table(['Option', 'Effet sur le binaire / assembleur', 'Conséquence et limites'], [
    ('-fstack-protector', 'Ajoute les opérations de sauvegarde et de vérification du canari dans les fonctions protégées. Séquence visible dans retlib-protected.', 'Arrête un retour après corruption détectée. La copie dangereuse et certaines corruptions restent possibles avant le contrôle.'),
    ('-fno-stack-protector', 'Supprime cette instrumentation pour la compilation concernée. bof du stack fourni ne montre pas de canari.', 'Facilite l’écrasement du cadre sans ce contrôle; ne désactive pas ASLR ni NX à lui seul.'),
    ('-z execstack / -z noexecstack', 'Modifie les permissions demandées pour la pile dans l’en-tête ELF; ne requiert pas une modification du corps de bof.', 'execstack autorise l’exécution sur pile; noexecstack la bloque. Return-to-libc réutilise du code exécutable existant.'),
    ('-m32', 'Produit un binaire et une convention d’appel 32 bits. Registres EBP/ESP et pointeurs de 4 octets visibles.', 'Les adresses sont encodées sur 4 octets. Les paramètres et offsets ne peuvent pas être réutilisés tels quels en 64 bits.'),
    ('-static', 'Lie les bibliothèques au binaire. file indique « statically linked » et le relevé ELF ne montre pas de segment dynamique.', 'Diffère du retlib observé avec libc dynamique; ses adresses system/exit ne peuvent pas être transplantées sans vérification.'),
])
p('La commande objdump -p stack indique un segment STACK avec les permissions rwx : la pile du fichier fourni est demandée exécutable. La capture E9 partie 2 montre bien -z execstack. Il ne faut donc pas présenter ce fichier comme une démonstration de return-to-libc avec NX actif. Aucune comparaison contrôlée avec -O2 n’a été fournie; son effet ne fait pas l’objet d’un résultat expérimental dans ce rapport.')

h('5. Défense par le typage et propriétés (E7, E8)')
h('5.1 Routine Ada et résultat observé (T7)', 2)
p('safe_input.adb définit Buffer_Type comme String(1..100). Il crée une constante Malicious_Input de 150 caractères A et tente de l’affecter à User_Buffer. En Ada, l’affectation entre ces tableaux de longueurs différentes échoue au contrôle de contrainte, sans écrire une chaîne de 150 caractères dans un objet de 100 caractères.')
code("subtype Buffer_Type is String (1 .. 100);\nUser_Buffer : Buffer_Type;\n-- Entrée de 150 caractères :\nUser_Buffer := Malicious_Input;\n-- Gestion au niveau de la procédure :\nexception\n   when Constraint_Error =>\n      Put_Line(\"[BLOQUÉ] Erreur de contrainte interceptée (Constraint_Error) !\");")
figure('lab2e7.png', 'Compilation GNAT et traitement de Constraint_Error')
p('La compilation avec gnatmake aboutit, avec des avertissements annonçant que 100 éléments sont attendus et que 150 sont fournis, puis que Constraint_Error sera levée à l’exécution. La trace d’exécution affiche le message [BLOQUÉ] et termine la procédure par son gestionnaire d’exception. Le résultat documente le rejet de l’affectation incompatible. Le message affiché sur l’absence de corruption est cohérent avec le contrôle de langage, sans constituer une mesure instrumentée de toute la mémoire.')
p('Le traitement aboutit à un état sûr pour ce programme : l’entrée n’est pas utilisée pour poursuivre un traitement et la procédure se termine après le diagnostic. Le tampon n’est pas réutilisé après l’exception. Cependant, cette entrée de 150 A n’est pas le badfile de 517 octets de T3 : le rejeu de la même entrée demandé par E7 reste à fournir. Le programme montre un sous-type de chaîne contraint; il n’utilise pas un sous-type d’index explicite distinct.')
h('5.2 Contrat présent dans check_contract.adb', 2)
code("subtype Buffer_Range is Natural range 1 .. 100;\nprocedure Process_Safe_Input (Data : String)\n  with Pre => Data'Length <= 100;")
figure('lab2e9.png', 'Compilation et appel valide de la procédure à précondition')
p('Le programme appelle Process_Safe_Input avec « Donnee conforme » et affiche « Entrée validée et conforme aux contrats de sécurité. ». Ce test présente une entrée valide; aucune entrée trop longue n’est appelée. Buffer_Range est déclaré mais n’est pas utilisé pour indexer un tampon, et la procédure ne réalise pas de copie. La trace démontre la compilation et l’exécution de l’appel valide, pas le rejet d’une précondition violée.')
p('La commande visible est gnatmake check_contract.adb, sans -gnata ni politique d’assertion explicite dans le source. La preuve d’un contrôle dynamique de la précondition n’est donc pas établie. Pour vérifier le contrat avec GNAT, il faut activer les assertions, par exemple avec -gnata ou pragma Assertion_Policy(Check), puis tracer un appel valide et un appel invalide. Le défaut de précondition active produit Assertion_Error; l’affectation de tableaux de tailles différentes produit Constraint_Error. Les deux mécanismes doivent être distingués.')
h('5.3 Propriétés explicites préparant les contrats SPARK', 2)
table(['Propriété', 'Formulation vérifiable', 'Portée dans le code fourni'], [
    ('P1 — longueur C', 'strlen(str) + 1 <= 100, avec un terminal nul présent dans les octets accessibles.', 'Non vérifiée avant strcpy; hypothèse violée par les charges trop longues.'),
    ('P2 — capacité de l’entrée Ada', "Data'Length <= 100 pour une entrée variable copiée dans un tampon de capacité 100.", 'Précondition écrite dans check_contract; activation et rejet non démontrés.'),
    ('P3 — longueur de l’affectation Ada actuelle', "Malicious_Input'Length = User_Buffer'Length pour l’affectation directe de tableaux.", '150 /= 100 : contrôle de contrainte levé et intercepté.'),
    ('P4 — borne d’index', "Pour toute écriture : I in User_Buffer'Range. Index normalisé dans 1..100.", 'Propriété proposée; aucune boucle de copie indexée dans les sources fournis.'),
    ('P5 — invariant après acceptation', 'Used_Length <= 100 et chaque caractère accepté est stocké à un index valide.', 'Propriété proposée pour une routine complète; Used_Length absent du code actuel.'),
    ('P6 — rejet sûr', 'Une entrée rejetée n’est ni exécutée ni transmise au traitement suivant; le programme termine ou conserve un état valide.', 'Illustrée par la terminaison du gestionnaire de safe_input.'),
    ('P7 — intégrité du contrôle', 'Adresse_retour_après = Adresse_retour_avant pour la fonction de copie.', 'Exigence de sécurité conceptuelle; aucune assertion exécutable de ce type dans le C fourni.'),
])
p('Une String(1..100) a une longueur exacte de 100, et non une longueur variable plafonnée à 100. Une affectation directe exige cette longueur exacte, même pour une entrée courte. Pour accepter des tailles de 0 à 100, une routine complète doit suivre la longueur utile et copier seulement la tranche admissible, ou utiliser une abstraction de chaîne bornée. Une précondition de longueur seule ne prouve pas l’implémentation; une preuve SPARK ultérieure devra relier contrat, indices et état du tampon.')

h('6. Corpus légué et limites')
p('L’énoncé exige au moins huit entrées documentées et rejouables dans corpus/. Aucun dossier corpus/ ni fichier badfile n’est fourni. Le tableau suivant définit huit cas à constituer; il ne les présente pas comme un corpus déjà exécuté. A^n désigne n octets 0x41, et 00 un terminal nul.')
table(['Cas / contenu à archiver', 'Objectif et propriété', 'État'], [
    ('C01 — A^100 suivi de 00 (101 octets)', 'Dépasse d’un octet la capacité C de 100 en comptant le terminal. P1.', 'Cas proposé, non rejoué.'),
    ('C02 — A^101 suivi de 00 (102 octets)', 'Dépassement explicite de la capacité de copie. P1/P4.', 'Cas proposé, non rejoué.'),
    ('C03 — A^150 suivi de 00 (151 octets)', 'Cas C correspondant à la longueur utile de l’essai Ada. P1/P3.', '150 caractères testés en Ada seulement; fichier C absent.'),
    ('C04 — A^200 suivi de 00 (201 octets)', 'Entrée longue pour la plage de tampons du Niveau 2. P1/P4.', 'Cas proposé; ne démontre pas une charge N2 réussie.'),
    ('C05 — 517 octets 0x41 sans nul', 'Teste aussi l’absence de terminaison après fread. P1.', 'Cas proposé, non rejoué.'),
    ('C06 — charge shellcode de la capture lab2e2.png (517 octets)', 'Écrasement de la cible de retour et code injecté. P7.', 'Génération et crash visibles; badfile original non fourni.'),
    ('C07 — charge return-to-libc du exploit.py fourni (300 octets)', 'Substitution de la cible par system et argument /bin/sh. P7.', 'Script présent; fichier généré et résultat non fournis.'),
    ('C08 — variante de C07 sans destination de retour exit valide', 'Étudie le retour après system et la disponibilité du processus. P7.', 'Variation demandée par E5, non exécutée.'),
])
p('Pour constituer le corpus effectif, chaque cas doit être conservé avec ses octets exacts, une description, le programme destinataire, les réglages de compilation et d’ASLR, le résultat attendu, le résultat observé et la trace. Les cas C06 à C08 nécessitent les adresses et offsets propres au binaire testé. Le tableau documentaire seul ne satisfait pas le livrable d’entrées rejouables.')
p('Le typage et les contrôles Ada limitent les écritures hors bornes lorsqu’ils sont conservés et que les opérations restent dans le modèle sûr du langage. Ils ne valident pas le sens des commandes, ne garantissent pas une absence de toute erreur logique et ne remplacent pas une preuve SPARK. Les protections mémoire compliquent certaines exploitations, mais ne corrigent pas strcpy. ASLR agit sur la prévisibilité des adresses, NX sur l’exécution de la pile, le canari sur certaines corruptions de cadre et dash sur la conservation des privilèges. Leurs garanties se complètent.')

h('7. Extension à distance — version Serveur (E9)')
figure('lab2e9part2.png', 'Compilation statique 32 bits et identification du binaire stack')
code('gcc -m32 -static -fno-stack-protector -z execstack stack.c -o stack\nfile stack')
p('La capture confirme un ELF 32 bits Intel 80386 lié statiquement. Cette compilation réalise une partie de la préparation technique d’E9. Le source stack.c fourni lit un fichier local et ne contient aucun traitement de socket. Dans SEED, ce type de programme peut être appelé par une infrastructure serveur distincte; aucun fichier de cette infrastructure n’est présent ici.')
p('Ni make install, ni dcbuild/dcup, ni docker-compose.yml, ni connexion TCP au port 9090 ne figurent dans les preuves. Aucun shell inversé ni shell obtenu à distance n’est montré. Le verdict T8 est donc : préparation de compilation documentée, exploitation distante non démontrée. La capture lab2e9.png, malgré son nom, documente E8 et non une attaque réseau.')
table(['Aspect', 'Attaque locale', 'Version Serveur attendue'], [
    ('Entrée', 'badfile ouvert par le processus local.', 'Charge envoyée via une connexion TCP au port 9090.'),
    ('Accès aux détails', 'Possibilité de désassembler le binaire et de relever la pile sous GDB.', 'Accès au binaire et aux adresses non garanti à l’attaquant distant.'),
    ('Adresses', 'Dépendent du binaire et de l’environnement local.', 'Dépendent du processus dans le conteneur; ne se déduisent pas directement des adresses locales.'),
    ('Interactivité', 'Shell attaché au contexte d’exécution local, si l’attaque réussit.', 'Shell distant ou inversé avec canal réseau établi, à démontrer.'),
    ('Preuve de réussite', 'Trace du shell et de son identité effective.', 'Trace client, trace conteneur, connexion et identité effective du shell distant.'),
])

h('8. Travail d’équipe')
p('Membres : Amel et Ibrahim Mat. Les noms ont été indiqués par Amel. Les captures montrent un terminal amoum@Amel, mais ne permettent pas d’attribuer à elles seules la conception, la réalisation ou la validation des parties à chacun.')
table(['Élément demandé', 'Information disponible / à compléter'], [
    ('Réalisation par Amel', 'À renseigner selon les contributions réelles.'),
    ('Réalisation par Ibrahim Mat', 'À renseigner selon les contributions réelles.'),
    ('Validation et rotation des rôles', 'À documenter pour chaque séance, avec l’inversion par rapport au laboratoire 1.'),
    ('Décisions et désaccords', 'Le contenu des essais est analysé dans ce rapport; les décisions d’équipe et désaccords éventuels ne sont pas documentés dans les pièces.'),
    ('Historique et dépôt Git', 'Adresse et références de commits non fournies.'),
    ('JOURNAL.md', 'Absent du dossier fourni; à joindre avec les dates, rôles, décisions et essais réels.'),
])
p('Les séances prévues par l’énoncé sont les 6 et 13 octobre 2026. Ce document, préparé le 5 octobre à partir des fichiers disponibles, ne prétend pas établir la présence ni le déroulement de ces séances. Les contributions doivent être complétées à partir du journal et de l’historique réel, sans reconstruire artificiellement des activités.')

h('9. Sources et déclaration d’usage de l’IA')
p('[1] Énoncé du laboratoire : CEG4799_CSI4539A_Lab2_Automne_2026.pdf, version 1.1, 3 octobre 2026, particulièrement les sections 3, 6, 7.1 et 10. Source des exigences, de la structure et des critères.')
p('[2] Fichiers fournis par l’équipe : stack.c, exploit.py, safe_input.adb, check_contract.adb et le binaire stack. Source de l’analyse du code et du relevé objdump.')
p('[3] Captures fournies : lab2e1.png à lab2e7.png, lab2e9.png et lab2e9part2.png. Source exclusive des résultats d’exécution historiques présentés.')
p('Les ressources SEED Labs et le manuel de Wenliang Du sont cités dans l’énoncé comme origine des activités. Leur consultation personnelle par les membres n’a pas été documentée dans les pièces fournies; toute source effectivement réutilisée par l’équipe doit être ajoutée à cette déclaration.')
p('Déclaration d’IA pour la préparation de ce document : OpenAI Codex, modèle GPT-6. Usage : lecture de l’énoncé, inspection des captures et des fichiers, désassemblage en lecture seule du binaire fourni, rédaction et mise en forme du rapport Word. Aucun résultat d’attaque réussi n’a été généré ou supposé. Les programmes fournis n’ont pas été modifiés. Les membres doivent relire le document, valider les informations et ajouter tout autre outil d’IA utilisé pour le code ou les essais.')

doc.add_page_break()
h('Annexe A. Tableau des essais T1 à T8')
p('« Partiel » signifie qu’une partie de la procédure est documentée sans atteindre tous les critères. Les résultats attendus ci-dessous proviennent de l’énoncé; les observations proviennent des pièces fournies.')
table(['Essai / procédure', 'Observation et preuve', 'Verdict'], [
    ('T1 — GDB : tampon, EBP, retour et distance.', 'lab2e1 : SIGSEGV, 0x90909090 et &buffer. Analyse du stack fourni : buffer EBP−108, retour EBP+4, distance 112.', 'Partiel : distance du fichier fourni établie; relevé GDB complet de stack-L1 et privilèges absents.'),
    ('T2 — Assemblage et tests autonomes 32/64 bits.', 'lab2e2 : 23 octets de shellcode 32 bits sans 00. Aucun a32.out/a64.out ni source assembleur fourni.', 'Partiel : code 32 bits analysé; tests autonomes et 64 bits non démontrés.'),
    ('T3 — Charge et débordement Niveaux 1 à 3.', 'lab2e3 : badfile de 517 octets; ./stack-L1 : Segmentation fault.', 'Échec du shell root en N1; N2 et N3 non démontrés.'),
    ('T4 — dash, ASLR, StackGuard et NX.', 'lab2e4 : ASLR=2; brute-force.sh absent; stack smashing detected après compilation -m32 -z noexecstack.', 'Partiel : canari détecté; contournements et NX isolé non démontrés.'),
    ('T5 — Retour vers system, NX actif.', 'lab2e5 : adresses system/exit. exploit.py : charge return-to-libc; adresse de chaîne supposée.', 'Partiel : préparation documentée; exécution, NX et identité effective absents.'),
    ('T6 — Source, assembleur, machine; options.', 'lab2e6 : canari dans retlib-protected. Relevé du stack fourni et octets analysés dans le rapport.', 'Partiel : chaîne détaillée pour stack et shellcode 32 bits; sources retlib et essais comparatifs contrôlés manquants.'),
    ('T7 — Rejeu E3 dans la défense Ada.', 'lab2e7 : compilation avec avertissements; affectation de 150 A rejetée par Constraint_Error géré.', 'Rejet Ada démontré pour 150 A; partiel pour le critère du rejeu exact de E3.'),
    ('T8 — Serveur en conteneur et TCP 9090.', 'lab2e9part2 : compilation -m32 -static -fno-stack-protector -z execstack; file confirme ELF i386 statique.', 'Préparation seulement; exploitation distante non démontrée.'),
])
p('E8 n’a pas d’essai distinct dans le tableau officiel. Le programme check_contract et sa trace documentent un appel valide; le corpus de huit fichiers et les essais de rejet du contrat restent à fournir.')

doc.add_page_break()
h('Annexe B. Code fourni et relevés du binaire')
p('Les sources ci-dessous sont reproduites telles que fournies, en UTF-8. Les commentaires contenant des valeurs supposées sont conservés pour traçabilité; leur présence ne valide pas ces valeurs. Les points techniques discutés dans le corps du rapport priment sur ces commentaires.')
for filename in ('stack.c', 'exploit.py', 'safe_input.adb', 'check_contract.adb'):
    h(filename, 2)
    code((ROOT / filename).read_text(encoding='utf-8'))
h('Relevé supplémentaire : bof du stack fourni', 2)
asm = subprocess.run(['objdump', '-Mintel', '-d', '--disassemble=bof', str(ROOT / 'stack')], capture_output=True, text=True, check=True).stdout
asm = asm.replace(str(ROOT / 'stack'), 'stack')
code(asm)
h('Relevé supplémentaire : permissions de la pile', 2)
headers = subprocess.run(['objdump', '-p', str(ROOT / 'stack')], capture_output=True, text=True, check=True).stdout
lines = headers.splitlines()
for i, line in enumerate(lines):
    if 'STACK' in line:
        code('\n'.join(lines[i:i+2]))
p('Les deux relevés objdump ont été obtenus pendant la rédaction. Aucun binaire vulnérable ni script d’attaque n’a été exécuté sur le poste de préparation du rapport.')

h('Annexe C. Éléments à compléter pour la remise')
table(['Exigence', 'Complément nécessaire'], [
    ('E1', 'Trace GDB complète du binaire testé, offset vérifié, compilation et permissions Set-UID documentées.'),
    ('E2', 'Sources assembleur et code machine 32/64 bits, outils d’extraction, traces des tests autonomes et absence de 00.'),
    ('E3', 'Scripts propres aux Niveaux 1, 2 et 3, traces de réussite avec id/getresuid et réglages de protections.'),
    ('E4', 'Essais dash avec setuid(0), force brute ASLR, comparaison du canari et essai NX isolé.'),
    ('E5', 'retlib.c, charge et offset validés sur retlib, chaîne d’environnement/prtenv, preuve NX, succès et variations exit/nom.'),
    ('E6', 'Source retlib correspondant à la capture et comparaisons contrôlées des options sur un même programme.'),
    ('E7', 'Rejeu des octets exacts de T3, sous-type d’index utilisé, projet .gpr et trace de rejet; avertissements à traiter.'),
    ('E8', 'Corpus d’au moins huit entrées effectives documentées; activation et essais des assertions.'),
    ('E9', 'server-code, attack-code, docker-compose.yml, installation/démarrage des conteneurs et shell distant tracé.'),
    ('Équipe et remise', 'Contributions réelles, JOURNAL.md, adresse du dépôt et historique; export du rapport en PDF et archive conforme.'),
])
p('Nom d’archive prévu par l’énoncé : CEG4799_Lab2_<NomA>_<NomB>.zip. Les noms de famille à utiliser dans cette convention doivent être confirmés par l’équipe. Ce rapport Word ne remplace pas les fichiers de code, le corpus et les traces demandés dans l’archive.')
doc.core_properties.title = 'Laboratoire 2 — Amel et Ibrahim Mat'
doc.core_properties.author = 'Amel et Ibrahim Mat'
doc.core_properties.subject = 'Attaques mémoire et défense par le typage — CEG4799 / CSI4539'
doc.save(OUT)
print(OUT)
print(f'Paragraphes : {len(doc.paragraphs)}; tableaux : {len(doc.tables)}; figures : {len(doc.inline_shapes)}')
